"""CV fictício em PDF (fpdf2, só latin-1) e DOCX, gerado no teste (sem binário no repo)."""

import io

from docx import Document
from fpdf import FPDF

LINHAS = [
    "Ana Exemplo - Cientista de Dados",
    "Experiência",
    "Gerente de Dados - Empresa Exemplo (nov/2022 - atual)",
    "Reduziu o custo da plataforma em 30% com ação contínua.",
]


def gerar_cv_pdf(linhas=LINHAS, paginas=1) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for n in range(paginas):
        pdf.add_page()
        for linha in linhas:
            pdf.cell(
                0, 8, linha if n == 0 else f"{linha} (p{n + 1})", new_x="LMARGIN", new_y="NEXT"
            )
    return bytes(pdf.output())


def gerar_cv_docx(linhas=LINHAS, tabela=None) -> bytes:
    doc = Document()
    for linha in linhas:
        doc.add_paragraph(linha)
    if tabela:
        t = doc.add_table(rows=1, cols=len(tabela))
        for celula, texto in zip(t.rows[0].cells, tabela, strict=True):
            celula.text = texto
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
