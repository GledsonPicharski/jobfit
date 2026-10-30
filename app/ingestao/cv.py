"""CV antigo em PDF ou DOCX: só o texto, em código. Os fatos vêm do `extrair_fatos`.

Um CV não tem layout previsível como o PDF do LinkedIn; por isso o código não tenta
achar cargos e datas aqui. Nos conflitos, o CV é a fonte de menor peso.
"""

import io
import re
import zipfile

FONTE = "cv"

_ASSINATURA_PDF = b"%PDF"
_ASSINATURA_ZIP = b"PK\x03\x04"
_ASSINATURA_DOC = b"\xd0\xcf\x11\xe0"  # Word 97-2003 (.doc)


class ErroCV(ValueError):
    pass


def _limpar(texto: str) -> str:
    linhas = [linha.rstrip() for linha in texto.replace("\r\n", "\n").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(linhas)).strip()


def _texto_do_pdf(conteudo: bytes) -> str:
    import pdfplumber  # lento para importar; só quando há PDF para ler
    from pdfplumber.utils.exceptions import PdfminerException

    try:
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            paginas = [pagina.extract_text() or "" for pagina in pdf.pages]
    except PdfminerException as erro:
        raise ErroCV("O arquivo não é um PDF que dê para ler.") from erro
    texto = _limpar("\n\n".join(paginas))
    if not texto:
        raise ErroCV(
            "O PDF não tem texto selecionável (deve ser uma imagem escaneada). "
            "Exporte o CV de novo a partir do Word ou do Google Docs."
        )
    return texto


def _texto_do_docx(conteudo: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = Document(io.BytesIO(conteudo))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as erro:
        raise ErroCV("O arquivo precisa ser um PDF ou DOCX.") from erro
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = list(dict.fromkeys(c.text.strip() for c in linha.cells))  # mescladas
            partes.append(" | ".join(c for c in celulas if c))
    return _limpar("\n".join(partes))


def texto_do_cv(conteudo: bytes) -> str:
    """O texto do CV, pelo conteúdo (não pela extensão, que pode estar errada)."""
    if conteudo.startswith(_ASSINATURA_PDF):
        return _texto_do_pdf(conteudo)
    if conteudo.startswith(_ASSINATURA_DOC):
        raise ErroCV("Arquivo .doc antigo: salve como .docx ou PDF no Word e envie de novo.")
    if conteudo.startswith(_ASSINATURA_ZIP):
        texto = _texto_do_docx(conteudo)
        if not texto:
            raise ErroCV("O DOCX está sem texto.")
        return texto
    raise ErroCV("O arquivo precisa ser um PDF ou DOCX.")
